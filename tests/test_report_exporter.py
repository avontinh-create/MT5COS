import base64
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.services.mt5_engine.report_exporter import MT5ReportExporter


class MT5ReportExporterTests(unittest.TestCase):
    def test_exports_report_tables_and_embeds_local_images(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            report = root / "report.htm"
            image = root / "chart.png"
            workbook_path = root / "report.xlsx"
            image.write_bytes(
                base64.b64decode(
                    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8"
                    "/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
                )
            )
            report.write_text(
                "<html><body>"
                "<table><tr><th>Metric</th><th>Value</th></tr>"
                "<tr><td>Total Net Profit:</td><td>25.5</td></tr></table>"
                '<img src="chart.png">'
                "</body></html>",
                encoding="utf-8",
            )

            result_path = MT5ReportExporter().export(
                str(report), str(workbook_path)
            )

            self.assertEqual(result_path, str(workbook_path.resolve()))
            self.assertTrue(workbook_path.is_file())
            workbook = load_workbook(workbook_path)
            self.assertEqual(workbook.sheetnames, ["MT5 Report", "Charts"])
            self.assertEqual(workbook["MT5 Report"]["A2"].value, "Total Net Profit:")
            self.assertEqual(len(workbook["Charts"]._images), 1)

    def test_fails_when_a_referenced_image_is_missing(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            report = root / "report.htm"
            report.write_text(
                '<html><body><img src="missing.png"></body></html>',
                encoding="utf-8",
            )

            with self.assertRaisesRegex(FileNotFoundError, "image"):
                MT5ReportExporter().export(
                    str(report), str(root / "report.xlsx")
                )


if __name__ == "__main__":
    unittest.main()
