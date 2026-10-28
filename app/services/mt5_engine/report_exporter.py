import os
import tempfile
from pathlib import Path
from urllib.parse import unquote, urlsplit

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.styles import Font


class MT5ReportExporter:
    @staticmethod
    def _read_html(report_path: Path) -> str:
        content = report_path.read_bytes()
        if content.startswith((b"\xff\xfe", b"\xfe\xff")):
            return content.decode("utf-16")
        if content.startswith(b"\xef\xbb\xbf"):
            return content.decode("utf-8-sig")
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("cp1252")

    @staticmethod
    def _local_image_paths(report_path: Path, html: str) -> list[Path]:
        soup = BeautifulSoup(html, "lxml")
        report_directory = report_path.parent.resolve()
        images = []
        seen = set()

        for image_tag in soup.find_all("img", src=True):
            source = urlsplit(image_tag["src"])
            if source.scheme or source.netloc or not source.path:
                continue

            image_path = (report_directory / unquote(source.path)).resolve()
            if os.path.commonpath((str(report_directory), str(image_path))) != str(
                report_directory
            ):
                continue
            if not image_path.is_file():
                raise FileNotFoundError(
                    f"MT5 report image was not found: {image_path}"
                )
            if image_path not in seen:
                images.append(image_path)
                seen.add(image_path)

        return images

    def export(self, report_path: str, workbook_path: str) -> str:
        source_path = Path(report_path).resolve()
        destination_path = Path(workbook_path).resolve()
        if not source_path.is_file():
            raise FileNotFoundError(f"MT5 report not found: {source_path}")
        if destination_path.suffix.lower() != ".xlsx":
            raise ValueError("Excel report output must use the .xlsx extension.")

        html = self._read_html(source_path)
        soup = BeautifulSoup(html, "lxml")
        image_paths = self._local_image_paths(source_path, html)

        workbook = Workbook()
        report_sheet = workbook.active
        report_sheet.title = "MT5 Report"
        report_sheet.freeze_panes = "A1"

        for table in soup.find_all("table"):
            for row in table.find_all("tr"):
                cells = row.find_all(["th", "td"], recursive=False)
                if not cells:
                    continue
                row_values = [
                    " ".join(cell.get_text(" ", strip=True).split())
                    for cell in cells
                ]
                report_sheet.append(row_values)
                for cell in report_sheet[report_sheet.max_row]:
                    cell.data_type = "s"
                    if any(element.name == "th" for element in cells):
                        cell.font = Font(bold=True)
            report_sheet.append([])

        if report_sheet.max_row == 1 and report_sheet["A1"].value is None:
            report_sheet["A1"] = "No tabular data was found in the MT5 report."

        for column_cells in report_sheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(
                (
                    len(str(cell.value))
                    for cell in column_cells
                    if cell.value is not None
                ),
                default=0,
            )
            report_sheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12), 60
            )

        if image_paths:
            chart_sheet = workbook.create_sheet("Charts")
            chart_sheet.column_dimensions["A"].width = 110
            next_row = 1
            for image_path in image_paths:
                chart_sheet.cell(next_row, 1, image_path.name).font = Font(bold=True)
                chart = SpreadsheetImage(str(image_path))
                chart.anchor = f"A{next_row + 1}"
                chart_sheet.add_image(chart)
                next_row += max(4, (chart.height + 19) // 20 + 3)

        destination_path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=destination_path.parent,
                prefix=f".{destination_path.stem}_",
                suffix=".xlsx",
                delete=False,
            ) as temporary_file:
                temporary_path = temporary_file.name
            workbook.save(temporary_path)
            os.replace(temporary_path, destination_path)
        finally:
            if temporary_path and os.path.exists(temporary_path):
                os.remove(temporary_path)

        return str(destination_path)
